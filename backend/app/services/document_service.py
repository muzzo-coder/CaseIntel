"""Document service handling file upload, validation, text extraction, and storage."""

import logging
import os
import shutil
import uuid
from pathlib import Path
from typing import List, Optional, Tuple
from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.core.config import settings
from app.core.errors import AppException, NotFoundException
from app.models.case import Case
from app.models.document import Document, DocumentStatus
from app.models.user import User
from app.services.case_service import get_case_by_id

logger = logging.getLogger("caseintel.services.document")

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md", ".rtf"}
MIME_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".rtf": "application/rtf",
}


def get_upload_dir() -> Path:
    """Ensure and return the base upload directory."""
    base_dir = Path(settings.UPLOAD_DIR)
    if not base_dir.is_absolute():
        base_dir = Path(__file__).resolve().parent.parent.parent / settings.UPLOAD_DIR
    base_dir.mkdir(parents=True, exist_ok=True)
    return base_dir


def extract_text_from_file(file_path: Path, ext: str) -> Tuple[str, int]:
    """Extract readable text and calculate page count from supported document formats."""
    text_parts: List[str] = []
    page_count = 1

    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            page_count = max(len(reader.pages), 1)
            for idx, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(f"[Page {idx + 1}]\n{extracted.strip()}")

        elif ext == ".docx":
            import docx
            doc = docx.Document(str(file_path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text_parts = paragraphs
            # Estimate pages roughly by word/paragraph length (approx 400 words per page)
            total_words = sum(len(p.split()) for p in paragraphs)
            page_count = max(1, (total_words // 400) + 1)

        else:  # Text, Markdown, plain files
            try:
                content = file_path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = file_path.read_text(encoding="latin-1", errors="replace")
            text_parts = [content.strip()]
            total_words = len(content.split())
            page_count = max(1, (total_words // 400) + 1)

    except Exception as e:
        logger.error(f"Text extraction failed for {file_path.name}: {e}")
        return f"[Text extraction failed: {str(e)}]", 1

    extracted_text = "\n\n".join(text_parts).strip()
    return extracted_text, page_count


def upload_case_document(
    db: Session,
    user: User,
    case_id: str,
    upload_file: UploadFile,
) -> Document:
    """Validate, store, extract text, and register a new document for a case workspace."""
    # Verify case ownership
    case = get_case_by_id(db, case_id, user)

    # Validate file extension
    orig_name = upload_file.filename or "uploaded_document.txt"
    ext = Path(orig_name).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise AppException(
            message=f"Unsupported file format '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            status_code=400,
            error_code="UNSUPPORTED_FILE_TYPE",
            details=[{"field": "file", "message": f"Must be one of {sorted(ALLOWED_EXTENSIONS)}"}],
        )

    # Save to disk
    doc_id = str(uuid.uuid4())
    case_dir = get_upload_dir() / "cases" / case_id
    case_dir.mkdir(parents=True, exist_ok=True)
    clean_filename = f"{doc_id}_{Path(orig_name).name}"
    dest_path = case_dir / clean_filename

    # Read and validate file size
    upload_file.file.seek(0, os.SEEK_END)
    file_size = upload_file.file.tell()
    upload_file.file.seek(0)

    max_bytes = settings.MAX_FILE_SIZE_MB * 1024 * 1024
    if file_size > max_bytes:
        raise AppException(
            message=f"File exceeds maximum size of {settings.MAX_FILE_SIZE_MB}MB.",
            status_code=400,
            error_code="FILE_TOO_LARGE",
        )

    with dest_path.open("wb") as buffer:
        shutil.copyfileobj(upload_file.file, buffer)

    # Extract text content
    extracted_text, page_count = extract_text_from_file(dest_path, ext)

    document = Document(
        id=doc_id,
        case_id=case.id,
        user_id=user.id,
        filename=orig_name,
        file_path=str(dest_path),
        file_size=file_size,
        mime_type=MIME_MAP.get(ext, upload_file.content_type or "application/octet-stream"),
        file_type=ext.lstrip("."),
        extracted_text=extracted_text,
        page_count=page_count,
        status=DocumentStatus.PROCESSED.value,
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    logger.info(f"Uploaded and extracted document ID={doc_id} ('{orig_name}') for case ID={case_id}")
    return document


def get_case_documents(db: Session, user: User, case_id: str) -> List[Document]:
    """Retrieve all documents belonging to a case workspace with ownership enforcement."""
    # Verify case ownership
    get_case_by_id(db, case_id, user)
    stmt = (
        select(Document)
        .where(Document.case_id == case_id, Document.user_id == user.id)
        .order_by(Document.created_at.desc())
    )
    return list(db.scalars(stmt).all())


def get_document_by_id(db: Session, user: User, document_id: str) -> Document:
    """Retrieve a single document owned by user."""
    stmt = select(Document).where(Document.id == document_id, Document.user_id == user.id)
    doc = db.scalars(stmt).first()
    if not doc:
        raise NotFoundException(message=f"Document with ID '{document_id}' was not found.")
    return doc


def delete_document(db: Session, user: User, document_id: str) -> None:
    """Permanently delete a document file and database entry."""
    doc = get_document_by_id(db, user, document_id)

    # Remove file from disk if present
    file_path = Path(doc.file_path)
    if file_path.exists():
        try:
            file_path.unlink()
        except Exception as e:
            logger.warning(f"Failed to remove file from disk {doc.file_path}: {e}")

    db.delete(doc)
    db.commit()
    logger.info(f"Deleted document ID={document_id} by user ID={user.id}")
